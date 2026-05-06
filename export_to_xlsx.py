"""Convert data_dump.json → AKD-회원정보.xlsx (2 sheets) + 2 TSVs.

Sheet 1 — 의사: id, 이름, 소속병원ID, 소속병원명, 직책, 전공, 한줄소개,
                  의사회직책, 경력, 학회회원, 강의, 수상, 진료분야, 링크
Sheet 2 — 병원: id, 병원명, 약식명, 주소, 전화번호, 위도, 경도, 홈페이지,
                  진료시간, 진료분야, 소속의사, 병원소개

배열/객체 필드는 줄바꿈으로 펼쳐서 한 셀에 표시.
사진 관련 필드(photo)는 제외(별도 처리).
"""
import json, csv, sys, io
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
data = json.loads((ROOT / 'data_dump.json').read_text(encoding='utf-8'))
HOSPITALS = data['HOSPITALS']
MEMBERS   = data['MEMBERS']

hosp_by_id = {h['id']: h for h in HOSPITALS}

# --- helpers ---
def join_lines(arr):
    return '\n'.join(arr) if arr else ''

def fmt_lectures(lectures):
    if not lectures: return ''
    return '\n'.join(f"{l.get('year','')} | {l.get('event','')} | {l.get('title','')}" for l in lectures)

def fmt_links(links):
    if not links: return ''
    return '\n'.join(f"{l.get('label','')} ({l.get('type','')}) | {l.get('url','')}" for l in links)

def fmt_hours(hours):
    if not hours: return ''
    return '\n'.join(f"{k}: {v}" for k, v in hours.items())

# ---------- 의사 시트 ----------
member_headers = [
    'id', '이름', '소속병원ID', '소속병원명', '직책', '전공', '한줄소개',
    '의사회 직책(역대)', '경력', '학회 회원', '강의 (연도|학술대회|제목)',
    '수상', '진료분야 (쉼표구분)', '링크 (이름|타입|URL)'
]
member_rows = []
for m in MEMBERS:
    h = hosp_by_id.get(m.get('hospitalId',''), {})
    member_rows.append([
        m.get('id',''),
        m.get('name',''),
        m.get('hospitalId',''),
        h.get('name',''),
        m.get('role',''),
        m.get('specialty',''),
        m.get('quote',''),
        join_lines(m.get('positions', [])),
        join_lines(m.get('career', [])),
        join_lines(m.get('memberships', [])),
        fmt_lectures(m.get('lectures', [])),
        join_lines(m.get('awards', [])),
        ', '.join(m.get('treatments', [])),
        fmt_links(m.get('links', [])),
    ])

# ---------- 병원 시트 ----------
hosp_headers = [
    'id', '병원명', '약식명', '주소', '전화번호', '위도', '경도', '홈페이지',
    '진료시간', '진료분야 (쉼표구분)', '소속 의사 ID', '소속 의사 이름', '병원 소개'
]
hosp_rows = []
mem_by_id = {m['id']: m for m in MEMBERS}
for h in HOSPITALS:
    doc_ids = h.get('doctorIds', [])
    doc_names = [mem_by_id.get(did, {}).get('name', did) for did in doc_ids]
    hosp_rows.append([
        h.get('id',''),
        h.get('name',''),
        h.get('shortName',''),
        h.get('address',''),
        h.get('phone',''),
        h.get('lat',''),
        h.get('lng',''),
        h.get('website',''),
        fmt_hours(h.get('hours', {})),
        ', '.join(h.get('treatments', [])),
        ', '.join(doc_ids),
        ', '.join(doc_names),
        h.get('description',''),
    ])

# ---------- XLSX ----------
wb = Workbook()
ws_m = wb.active
ws_m.title = '의사'
ws_h = wb.create_sheet('병원')

HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill('solid', fgColor='1B4965')
WRAP = Alignment(wrap_text=True, vertical='top')
BORDER = Border(left=Side(style='thin', color='D0D7DE'),
                right=Side(style='thin', color='D0D7DE'),
                top=Side(style='thin', color='D0D7DE'),
                bottom=Side(style='thin', color='D0D7DE'))

def write_sheet(ws, headers, rows, col_widths):
    ws.append(headers)
    for c, _ in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.alignment = WRAP
            cell.border = BORDER
    ws.row_dimensions[1].height = 28
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 130
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

# col widths tuned for korean wrap
member_widths = [14, 10, 18, 24, 8, 14, 36, 28, 28, 24, 38, 22, 36, 38]
hosp_widths   = [18, 26, 18, 36, 16, 8, 8, 26, 28, 38, 22, 22, 36]
write_sheet(ws_m, member_headers, member_rows, member_widths)
write_sheet(ws_h, hosp_headers,   hosp_rows,   hosp_widths)

xlsx_path = ROOT / 'AKD-회원정보.xlsx'
wb.save(xlsx_path)
print('saved:', xlsx_path)

# ---------- TSV (2개) ----------
def write_tsv(path, headers, rows):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, delimiter='\t', quoting=csv.QUOTE_MINIMAL)
        w.writerow(headers)
        for r in rows:
            # replace \n with literal \n for cleaner pasting into Sheets
            w.writerow([str(c).replace('\n', ' \\n ') if isinstance(c, str) else c for c in r])

write_tsv(ROOT / 'AKD-의사.tsv', member_headers, member_rows)
write_tsv(ROOT / 'AKD-병원.tsv', hosp_headers,   hosp_rows)
print('TSVs saved.')
